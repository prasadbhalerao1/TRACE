"""Format Normalization + Content Extraction — doc 04 §4, tool-only (no LLM).

`.pptx` -> python-pptx directly. `.pdf` -> PyMuPDF (text + rasterized page image, since
a PDF export of a deck has no separate "notes" concept — each page IS the slide).
Legacy `.ppt` needs a LibreOffice headless conversion to `.pptx` first; that binary
isn't bundled with this repo's Python venv, so it degrades via a typed error rather
than silently failing, same pattern as `ResumeExtractionUnavailable`.
"""

import io
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import fitz  # PyMuPDF
from pptx import Presentation as PptxPresentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from services.api.core.config import get_settings


class UnsupportedDeckFormat(ValueError):
    """Raised when the uploaded file isn't .pptx, .ppt, or .pdf."""


class LegacyPptConversionUnavailable(RuntimeError):
    """Raised when a .ppt was uploaded but LibreOffice headless isn't available to
    convert it — never silently skip content extraction for a legacy file."""


@dataclass
class ExtractedSlide:
    index: int
    title: str | None
    body: str | None
    notes: str | None
    has_image: bool
    image_bytes: list[bytes] = field(default_factory=list)


def detect_format(file_name: str | None, content_type: str | None) -> str:
    name = (file_name or "").lower()
    ctype = (content_type or "").lower()
    if name.endswith(".pptx") or "presentationml" in ctype:
        return "pptx"
    if name.endswith(".ppt") or ctype == "application/vnd.ms-powerpoint":
        return "ppt"
    if name.endswith(".pdf") or ctype == "application/pdf":
        return "pdf"
    raise UnsupportedDeckFormat(f"Unsupported deck format: name={file_name!r} content_type={content_type!r}")


def convert_ppt_to_pptx(file_bytes: bytes, binary: str = "soffice") -> bytes:
    """Legacy .ppt -> .pptx via `soffice --headless --convert-to pptx`."""
    resolved = shutil.which(binary)
    if resolved is None:
        raise LegacyPptConversionUnavailable(
            f"LibreOffice binary '{binary}' not found on PATH — cannot convert legacy .ppt."
        )

    with tempfile.TemporaryDirectory() as tmpdir:
        src = Path(tmpdir) / "input.ppt"
        src.write_bytes(file_bytes)
        try:
            subprocess.run(
                [resolved, "--headless", "--convert-to", "pptx", "--outdir", tmpdir, str(src)],
                check=True,
                capture_output=True,
                timeout=get_settings().libreoffice_timeout_seconds,
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
            raise LegacyPptConversionUnavailable(f"LibreOffice conversion failed: {exc}") from exc

        out = Path(tmpdir) / "input.pptx"
        if not out.exists():
            raise LegacyPptConversionUnavailable("LibreOffice did not produce a .pptx output file.")
        return out.read_bytes()


def extract_slides_from_pptx(file_bytes: bytes) -> list[ExtractedSlide]:
    deck = PptxPresentation(io.BytesIO(file_bytes))
    slides: list[ExtractedSlide] = []

    for i, slide in enumerate(deck.slides):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape and title_shape.text else None
        # python-pptx's `slide.shapes.title` re-wraps a fresh proxy object on every
        # access, so `shape is not slide.shapes.title` never actually excludes it —
        # compare by shape_id instead, which is stable across accesses.
        title_shape_id = title_shape.shape_id if title_shape is not None else None

        body_parts: list[str] = []
        image_bytes: list[bytes] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.shape_id != title_shape_id:
                text = shape.text_frame.text.strip()
                if text:
                    body_parts.append(text)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image_bytes.append(shape.image.blob)
                except Exception:
                    pass

        notes = None
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            notes = notes_text or None

        slides.append(
            ExtractedSlide(
                index=i,
                title=title,
                body="\n".join(body_parts) or None,
                notes=notes,
                has_image=bool(image_bytes),
                image_bytes=image_bytes,
            )
        )

    return slides


def extract_slides_from_pdf(file_bytes: bytes) -> list[ExtractedSlide]:
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    slides: list[ExtractedSlide] = []

    for i, page in enumerate(doc):
        text = page.get_text().strip()
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        title = lines[0] if lines else None
        body = "\n".join(lines[1:]) if len(lines) > 1 else None

        # A PDF deck export has no OCR/notes distinction — rasterize the whole page so
        # the Slide Image/OCR Agent can still run diagram/chart understanding on it.
        pixmap = page.get_pixmap(dpi=100)
        image_bytes = [pixmap.tobytes("png")]

        slides.append(
            ExtractedSlide(
                index=i,
                title=title,
                body=body,
                notes=None,
                has_image=True,
                image_bytes=image_bytes,
            )
        )

    return slides
